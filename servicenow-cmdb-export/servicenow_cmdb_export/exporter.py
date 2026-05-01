"""Excel writer for CMDB export rows."""
from __future__ import annotations

from pathlib import Path
from typing import Iterable

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from servicenow_cmdb_export.queries import CMDBSnapshot, Row


def write_workbook(
    path: str | Path,
    rows: Iterable[Row],
    snapshot: CMDBSnapshot | None = None,
    include_raw: bool = False,
) -> int:
    wb = Workbook()
    ws = wb.active
    ws.title = "device_app_map"

    headers = Row.headers()
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)

    count = 0
    for row in rows:
        ws.append(row.as_list())
        count += 1

    _autosize(ws, len(headers))
    ws.freeze_panes = "A2"

    if include_raw and snapshot is not None:
        _add_raw_sheet(wb, "raw_servers", snapshot.servers)
        _add_raw_sheet(wb, "raw_clients", snapshot.clients)
        _add_raw_sheet(wb, "raw_apps", list(snapshot.apps_by_id.values()))
        _add_raw_sheet(wb, "raw_envs", list(snapshot.envs_by_id.values()))

    wb.save(path)
    return count


def _autosize(ws, col_count: int, max_width: int = 50) -> None:
    for idx in range(1, col_count + 1):
        letter = get_column_letter(idx)
        longest = 0
        for cell in ws[letter]:
            v = cell.value
            if v is None:
                continue
            longest = max(longest, len(str(v)))
        ws.column_dimensions[letter].width = min(longest + 2, max_width)


def _add_raw_sheet(wb: Workbook, title: str, records: list[dict]) -> None:
    ws = wb.create_sheet(title=title)
    if not records:
        ws.append(["(no records)"])
        return
    keys = sorted({k for r in records for k in r.keys()})
    ws.append(keys)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for r in records:
        ws.append([_flatten(r.get(k)) for k in keys])
    _autosize(ws, len(keys))
    ws.freeze_panes = "A2"


def _flatten(field) -> str:
    if isinstance(field, dict):
        return field.get("display_value") or field.get("value") or ""
    if field is None:
        return ""
    return str(field)
